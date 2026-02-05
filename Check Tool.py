# pip install python-docx rapidfuzz openpyxl

import re
import io
import os
import base64
import pandas as pd
import numpy as np
import streamlit as st
from pathlib import Path
from docx import Document
from rapidfuzz import fuzz
from io import BytesIO


# ============================================================
# Streamlit page config (MUST be before any other st.* calls)
# ============================================================
st.set_page_config(
    page_title="Check Tool",
    page_icon="🔎",
    layout="centered"
)


# ============================================================
# Global CSS styling (button styles, info box text size, header)
# Put all CSS here (top of the page) so it is applied consistently
# ============================================================
st.markdown("""
<style>
/* ===== Orange Download Button ===== */
div.stDownloadButton > button {
    background-color: #f97316;
    color: white;
    font-size: 16px;
    font-weight: 600;
    padding: 0.65em 1.5em;
    border-radius: 12px;
    border: none;
    box-shadow: 0 4px 10px rgba(0,0,0,0.18);
    transition: all 0.2s ease-in-out;
}

div.stDownloadButton > button:hover {
    background-color: #ea580c;
    transform: translateY(-1px);
    box-shadow: 0 6px 14px rgba(0,0,0,0.25);
}

div.stDownloadButton > button:active {
    transform: scale(0.97);
}

/* ===== Smaller font inside st.info (alerts) ===== */
div[data-testid="stAlert"] {
    font-size: 12px !important;
    line-height: 1.5 !important;
    padding-top: 0.6rem;
    padding-bottom: 0.6rem;
}

div[data-testid="stAlert"] ul,
div[data-testid="stAlert"] li {
    font-size: 12px !important;
    line-height: 1.5 !important;
}

/* ===== Header (logo + title) ===== */
.gsk-header {
    display: flex;
    align-items: center;        /* vertical center */
    justify-content: center;    /* horizontal center */
    gap: 16px;
    margin: 18px 0 14px 0;
}

.gsk-header img {
    height: 38px;              /* logo height */
    width: auto;
    display: block;
}

.gsk-header .title {
    font-size: 44px;
    font-weight: 800;
    color: #2b2d38;
    margin: 0;
    line-height: 1;            /* remove extra vertical whitespace */
}
</style>
""", unsafe_allow_html=True)


# ============================================================
# Header: GSK logo + "Check Tool" title (centered)
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
LOGO_PATH = BASE_DIR / "GSK.png"

def img_to_base64(path: Path) -> str:
    """Read an image file and return base64 string so it can be embedded in HTML."""
    return base64.b64encode(path.read_bytes()).decode("utf-8")

logo_b64 = img_to_base64(LOGO_PATH) if LOGO_PATH.exists() else ""

st.markdown(f"""
<div class="gsk-header">
    {"<img src='data:image/png;base64," + logo_b64 + "' alt='GSK logo'/>" if logo_b64 else ""}
    <div class="title">Check Tool</div>
</div>
""", unsafe_allow_html=True)


# ============================================================
# User instructions (file naming + required Word table header)
# ============================================================
st.info("""
📄 **File naming**
- Filename must end with **WN##.docx**
- Example: *Cleaning Risk Assessment WN16.docx*

📊 **Table requirements inside the Word file**
- The document must contain **one table**
- The table header must include:
  - **Composé**
  - **CAS** or **CAS N°**
""")


# ============================================================
# Upload ONE Word file (.docx)
# ============================================================
uploaded_file = st.file_uploader(
    "Upload ONE Word (.docx) file",
    type=["docx"],
    accept_multiple_files=False
)


# ============================================================
# Session state (persist data across Streamlit reruns)
# ============================================================
if "wn_df" not in st.session_state:
    st.session_state["wn_df"] = None
if "filename" not in st.session_state:
    st.session_state["filename"] = None


# ============================================================
# Helper: Extract the target table from Word into a DataFrame
# - Find the table where the first 2 rows contain both 'Composé' and 'CAS'
# - Keep only the first 2 columns
# ============================================================
def extract_target_table_to_df(doc: Document) -> pd.DataFrame:
    """Find the table containing 'Composé' and 'CAS' and return the first 2 columns as DataFrame."""
    target_table = None

    for table in doc.tables:
        header_text = " ".join(
            cell.text.strip()
            for row in table.rows[:2]
            for cell in row.cells
        )
        if "Composé" in header_text and "CAS" in header_text:
            target_table = table
            break

    if target_table is None:
        raise ValueError(
            "No valid table found.\n"
            "Table header must include 'Composé' and 'CAS' / 'CAS N°'."
        )

    df = pd.DataFrame(
        [[cell.text.strip() for cell in row.cells] for row in target_table.rows]
    )

    # First row is the header
    df.columns = df.iloc[0]
    df = df.iloc[1:].reset_index(drop=True)

    # Keep only the first two columns (Composé, CAS...)
    df = df.iloc[:, :2]

    return df


# ============================================================
# Load Excel reference table (same folder as the script)
# ============================================================
excel_path = BASE_DIR / "TTCG04 pour analyse.xlsx"

# Using Path for excel_path, so use excel_path.exists() rather than os.path.exists(...)
if not excel_path.exists():
    st.error("❌ Excel file 'TTCG04 pour analyse.xlsx' not found in the app folder.")
    st.stop()

df_main = pd.read_excel(excel_path, usecols="A:B")
df_main.columns = df_main.iloc[1]
df_main = df_main.iloc[2:].reset_index(drop=True)
df_main = df_main.dropna(how="all")


# ============================================================
# Process Word upload
# - Stop early if no file uploaded
# - Extract target table and store into session_state
# ============================================================
if uploaded_file is None:
    st.info("Please upload one .docx file.")
    st.stop()

try:
    doc = Document(BytesIO(uploaded_file.getvalue()))
    wn_df = extract_target_table_to_df(doc)

    # Persist across reruns
    st.session_state["wn_df"] = wn_df
    st.session_state["filename"] = uploaded_file.name

    st.success(f"Loaded table from: {uploaded_file.name}")
    st.subheader("Preview extracted table")
    st.dataframe(wn_df, width="stretch")

except Exception as e:
    st.error(str(e))
    st.stop()


# ============================================================
# Retrieve data like in a notebook (safe guard)
# ============================================================
if st.session_state["wn_df"] is None:
    st.info("Please upload a valid Word file to continue.")
    st.stop()

# Copy to avoid mutating session_state data directly
wn_df = st.session_state["wn_df"].copy()
main = df_main.copy()


# ============================================================
# Rename columns to standard names
# ============================================================
wn_df = wn_df.rename(columns={"CAS N°": "cas_no", "Composé": "compose"})
main = main.rename(columns={"CAS N°": "cas_no", "Compounds / Ingredients": "compound"})


# ============================================================
# Clean strings and treat NQ/NA/etc. as missing CAS
# ============================================================
wn_df["cas_no"] = wn_df["cas_no"].astype(str).str.strip()
main["cas_no"] = main["cas_no"].astype(str).str.strip()
wn_df["compose"] = wn_df["compose"].astype(str).str.strip()
main["compound"] = main["compound"].astype(str).str.strip()

wn_df.loc[
    wn_df["cas_no"].str.upper().isin(["NQ", "NA", "N/A", "NONE", ""]),
    "cas_no"
] = pd.NA


# ============================================================
# Normalization helpers for name matching
# ============================================================
def norm_name(s):
    """Normalize a chemical/compound name for robust matching."""
    if pd.isna(s):
        return ""
    s = str(s).lower().strip()

    # Remove non-breaking spaces and zero-width spaces
    s = (s.replace("\xa0", " ")
           .replace("\u200b", " ")
           .replace("\u2009", " ")
           .replace("\u202f", " "))

    # Normalize different dash/hyphen characters
    s = re.sub(r"[\u2010\u2011\u2012\u2013\u2014\u2212\-]+", " ", s)

    # Remove parentheses and middle dots
    s = s.replace("(", " ").replace(")", " ")
    s = s.replace("·", " ")

    # Keep "/" for alias splitting; convert other symbols to spaces
    s = re.sub(r"[^a-z0-9/]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()

    # Merge patterns like "men a" -> "mena" (repeat a few times)
    for _ in range(3):
        s = re.sub(r"\b([a-z]{2,})\s+([a-z])\b", r"\1\2", s)

    return s

def split_aliases(s_norm):
    """Split 'a/b/c' into ['a','b','c'] after normalization."""
    if not s_norm:
        return []
    return [p.strip() for p in s_norm.split("/") if p.strip()]


# ============================================================
# Similarity scoring (RapidFuzz)
# - token_set_ratio: good for different token order / extra words
# - partial_ratio: good if one name contains the other
# ============================================================
def similarity_score(a_name, b_name):
    """Return similarity score (0-100) using multiple RapidFuzz scorers."""
    a = norm_name(a_name)
    b = norm_name(b_name)
    if not a or not b:
        return 0
    if a == b:
        return 100

    candidates = split_aliases(b) or [b]
    best = 0
    for c in candidates:
        best = max(best, int(max(
            fuzz.token_set_ratio(a, c),
            fuzz.partial_ratio(a, c),
        )))
    return best


# ============================================================
# Thresholds and token gating settings
# ============================================================
TH_CAS = 60      # threshold for "match" when CAS is present
TH_NO_CAS = 65   # threshold to consider "present" when CAS is missing

STOPWORDS = {
    "antigen", "conjugate", "strain", "virus", "toxoid", "purified", "live", "attenuated",
    "like", "killed", "protein", "type", "surface", "ads", "for", "from", "form"
}

def key_tokens(s):
    """Extract key tokens (>=3 chars) excluding stopwords for token-gated fuzzy matching."""
    s = norm_name(s)
    toks = [t for t in s.split() if len(t) >= 3 and t not in STOPWORDS]
    return set(toks)

MIN_OVERLAP = 2  # minimum shared key tokens before computing fuzzy similarity

def fuzzy_score(a_name, b_name):
    """Fuzzy score for NO-CAS matching (avoid too-permissive partial matching)."""
    a = norm_name(a_name)
    b = norm_name(b_name)
    if not a or not b:
        return 0
    if a == b:
        return 100
    return int(max(fuzz.WRatio(a, b), fuzz.token_set_ratio(a, b)))


# ============================================================
# 1) HAS CAS: Join by CAS -> present/absent -> validate name match
# ============================================================
wn_has_cas = wn_df[wn_df["cas_no"].notna()].copy()
main_small = main[["cas_no", "compound"]].copy()

wn_cas_m = wn_has_cas.merge(main_small, on="cas_no", how="left")
wn_cas_m["present"] = np.where(wn_cas_m["compound"].notna(), "present", "absent")

# Only compute name similarity when present
wn_cas_m["_score"] = np.where(
    wn_cas_m["present"] == "present",
    wn_cas_m.apply(lambda r: similarity_score(r["compose"], r["compound"]), axis=1),
    0
)

# If multiple main rows share same CAS, keep the best match per (cas_no, compose)
wn_cas_best = (
    wn_cas_m.sort_values(["cas_no", "_score"], ascending=[True, False])
            .drop_duplicates(subset=["cas_no", "compose"], keep="first")
            .reset_index(drop=True)
)

wn_cas_best["match"] = np.where(
    (wn_cas_best["present"] == "present") & (wn_cas_best["_score"] >= TH_CAS),
    "yes", "no"
)
wn_cas_best["group"] = "has_cas"


# ============================================================
# 2) NO CAS: Strict normalized match first, then token-gated fuzzy
# ============================================================
wn_no_cas = wn_df[wn_df["cas_no"].isna()].copy()
main_compounds = main["compound"].dropna().astype(str).tolist()

# Build normalized mapping for strict matching
main_norm_map = {}
for c in main_compounds:
    k_norm = norm_name(c)
    if k_norm and k_norm not in main_norm_map:
        main_norm_map[k_norm] = c

def best_match_in_main_no_cas(wn_name):
    """Return (best_compound, score, present/absent) for a WN name without CAS."""
    if pd.isna(wn_name) or not str(wn_name).strip():
        return ("", 0, "absent")

    a_norm = norm_name(wn_name)

    # Stage 1: strict exact match on normalized form
    if a_norm in main_norm_map:
        return (main_norm_map[a_norm], 100, "present")

    # Stage 2: token-gated fuzzy matching
    a_tokens = key_tokens(wn_name)
    if len(a_tokens) == 0:
        return ("", 0, "absent")

    best_c, best_s = "", 0
    for c in main_compounds:
        # Token gate to avoid everyone matching the same one
        if len(a_tokens & key_tokens(c)) < MIN_OVERLAP:
            continue

        s = fuzzy_score(wn_name, c)
        if s > best_s:
            best_s, best_c = s, c
            if best_s == 100:
                break

    present = "present" if best_s >= TH_NO_CAS else "absent"
    return (best_c, best_s, present)

best_triplets = wn_no_cas["compose"].apply(best_match_in_main_no_cas)

wn_no_cas["compound"] = best_triplets.apply(lambda x: x[0])
wn_no_cas["name_score"] = best_triplets.apply(lambda x: x[1])
wn_no_cas["present"] = best_triplets.apply(lambda x: x[2])
wn_no_cas["match"] = np.where(wn_no_cas["present"] == "present", "yes", "no")
wn_no_cas["group"] = "no_cas"


# ============================================================
# 3) Combine + finalize columns
# - If absent, clear compound to avoid misleading "looks matched" output
# ============================================================
wn_final = pd.concat([wn_cas_best, wn_no_cas], ignore_index=True)

wn_final.loc[wn_final["present"] == "absent", "compound"] = pd.NA

cols_keep = ["group", "cas_no", "compose", "present", "compound", "match"]
wn_final = wn_final[cols_keep].copy()


# ============================================================
# 4) Summary metrics (ratios)
# ============================================================
total = len(wn_final)
present_ratio = ((wn_final["present"] == "present").sum() / total) if total else 0
match_ratio = ((wn_final["match"] == "yes").sum() / total) if total else 0

st.markdown(
    """
    <div style="display: flex; align-items: center; gap: 6px;">
        <h3 style="margin: 0;">Summary</h3>
        <span
            title="Present ratio represents the proportion of items from the uploaded Word file that are found in the reference Excel dataset.
Match ratio represents the proportion of items that are both present and successfully matched based on CAS number and/or compound name similarity rules."
            style="cursor: help; font-size: 16px; color: #6b7280;"
        >ℹ️</span>
    </div>
    """,
    unsafe_allow_html=True
)

st.write(f"Present ratio: {present_ratio:.2%}")
st.write(f"Match ratio: {match_ratio:.2%}")

# If you want to display the full final table, uncomment below:
# st.subheader("Final result table")
# st.dataframe(wn_final, width="stretch")


# ============================================================
# 5) Export result to Excel in-memory + Download button
# Also write summary values into G1/H1
# ============================================================
G1_value = f"Present ratio: {present_ratio:.2%}"
H1_value = f"Match ratio: {match_ratio:.2%}"

output = io.BytesIO()
with pd.ExcelWriter(output, engine="openpyxl") as writer:
    wn_final.to_excel(writer, index=False, sheet_name="Sheet1")
    ws = writer.sheets["Sheet1"]
    ws["G1"] = G1_value
    ws["H1"] = H1_value

# Extract WN## (e.g., WN16) from uploaded filename for a clean output name
m = re.search(r"(WN\d{2})", uploaded_file.name, re.IGNORECASE)
wn_code = m.group(1).upper() if m else "WN_UNKNOWN"
out_filename = f"{wn_code}_check_result.xlsx"

st.download_button(
    label="⬇️ Download Excel result",
    data=output.getvalue(),
    file_name=out_filename,
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
